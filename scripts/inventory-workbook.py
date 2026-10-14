"""Read-only inventory. Requires openpyxl; never saves the source workbook.

Usage: python inventory-workbook.py reviewed.xlsx inventory.json [original.xlsx]
Shared formulas are expanded by the reader so followers are included.
"""
import hashlib
import json
from pathlib import Path
import sys
import openpyxl

path = Path(sys.argv[1])
formulas = openpyxl.load_workbook(path, data_only=False)
values = openpyxl.load_workbook(path, data_only=True)
result = {'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheets': []}
for sheet in formulas:
    cells = [c for row in sheet for c in row if c.value is not None]
    computed = [c for c in cells if c.data_type == 'f']
    result['sheets'].append({
        'name': sheet.title, 'state': sheet.sheet_state,
        'populatedCells': len(cells), 'formulaCells': len(computed),
        'cachedErrors': [c.coordinate for row in values[sheet.title] for c in row if c.data_type == 'e'],
        'missingFormulaCaches': [c.coordinate for c in computed if values[sheet.title][c.coordinate].value is None],
        'formulaInventory': {c.coordinate: c.value for c in computed},
    })
if len(sys.argv) > 3:
    original = Path(sys.argv[3])
    before = openpyxl.load_workbook(original, data_only=False)
    result['original'] = {'file': original.name, 'sha256': hashlib.sha256(original.read_bytes()).hexdigest(), 'sheetNames': before.sheetnames}
    result['changedCellsFromOriginal'] = [
        {'sheet': sheet.title, 'cell': c.coordinate, 'original': c.value, 'reviewed': formulas[sheet.title][c.coordinate].value}
        for sheet in before for row in sheet for c in row
        if c.value != formulas[sheet.title][c.coordinate].value
    ]
Path(sys.argv[2]).write_text(json.dumps(result, indent=2, ensure_ascii=False) + '\n')
assert hashlib.sha256(path.read_bytes()).hexdigest() == result['sha256']
print(json.dumps({'sheets': len(result['sheets']), 'formulaCells': sum(s['formulaCells'] for s in result['sheets']), 'sourceUnchanged': True}))
